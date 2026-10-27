"""Fixtures for the worker's tests. Every file is made here, on the fly: no binary fixture is committed."""
from __future__ import annotations

import io
import os
import secrets
import shutil
import subprocess
import sys
import zipfile
import zlib
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient  # noqa: E402

from app import security  # noqa: E402
from app.config import Settings  # noqa: E402
from app.main import create_app  # noqa: E402

SECRET = b"test-secret-0123456789abcdef0123456789"


class Signed:
    """A client that signs its requests the way the VPS does (src/media/worker.ts)."""

    def __init__(self, client: TestClient, secret: bytes = SECRET, clock=None) -> None:
        self.client = client
        self.secret = secret
        self.clock = clock or __import__("time").time

    def headers(self, method: str, path_with_query: str, body: bytes, *, ts: int | None = None, nonce: str | None = None, secret: bytes | None = None) -> dict:
        ts = int(self.clock()) if ts is None else ts
        nonce = nonce or secrets.token_hex(16)
        sig = security.sign(secret or self.secret, ts, nonce, method, path_with_query, body)
        return {"x-tailzu-timestamp": str(ts), "x-tailzu-nonce": nonce, "x-tailzu-signature": sig, "content-type": "application/octet-stream"}

    def post(self, path_with_query: str, body: bytes = b"", **kw):
        return self.client.post(path_with_query, content=body, headers=self.headers("POST", path_with_query, body, **kw))


@pytest.fixture
def settings() -> Settings:
    return Settings(secret=SECRET)


@pytest.fixture
def client(settings) -> Signed:
    return Signed(TestClient(create_app(settings)))


# ------------------------------------------------------------------ files made on the fly


def make_docx(paragraphs=("Hello world",), heading="Employee handbook") -> bytes:
    import docx

    d = docx.Document()
    d.add_heading(heading, level=1)
    for p in paragraphs:
        d.add_paragraph(p)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Item", "Days", "Notice", "60"
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def make_xlsx(sheets: dict | None = None) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in (sheets or {"Sales": [["Month", "Revenue"], ["Jan", 100], ["Feb", 150], ["Mar", 120]]}).items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def make_pptx(slides=(("Launch plan", "We ship in March."), ("Budget", "Four lakh rupees."))) -> bytes:
    from pptx import Presentation

    p = Presentation()
    for title, body in slides:
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
    out = io.BytesIO()
    p.save(out)
    return out.getvalue()


ODF_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"'
)


def make_odf(kind: str, body: str, encrypted: bool = False) -> bytes:
    mime = {"text": "application/vnd.oasis.opendocument.text", "spreadsheet": "application/vnd.oasis.opendocument.spreadsheet", "presentation": "application/vnd.oasis.opendocument.presentation"}[kind]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("mimetype", mime, compress_type=zipfile.ZIP_STORED)
        z.writestr("content.xml", f'<?xml version="1.0"?><office:document-content {ODF_NS}><office:body><office:{kind}>{body}</office:{kind}></office:body></office:document-content>')
        z.writestr("META-INF/manifest.xml", '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">' + ("<manifest:encryption-data/>" if encrypted else "") + "</manifest:manifest>")
    return out.getvalue()


def make_epub(drm: bool = False, font_obfuscation: bool = False) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        z.writestr("OEBPS/content.opf", '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>A Book</dc:title><dc:identifier id="id">x</dc:identifier><dc:language>en</dc:language></metadata><manifest><item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="c1"/></spine></package>')
        z.writestr("OEBPS/ch1.xhtml", '<html xmlns="http://www.w3.org/1999/xhtml"><body><h1>Chapter One</h1><p>It was a bright cold day in April.</p></body></html>')
        if drm or font_obfuscation:
            algo = "http://www.w3.org/2001/04/xmlenc#aes128-cbc" if drm else "http://www.idpf.org/2008/embedding"
            z.writestr("META-INF/encryption.xml", f'<encryption xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><EncryptedData><EncryptionMethod Algorithm="{algo}"/></EncryptedData></encryption>')
    return out.getvalue()


def make_pdf(pages=("Refund policy: thirty days.",), encrypt: str | None = None) -> bytes:
    objs: dict[int, str] = {1: "<< /Type /Catalog /Pages 2 0 R >>", 3: "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"}
    kids, n = [], 4
    for text in pages:
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objs[n] = f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {n + 1} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        objs[n + 1] = f"<< /Length {len(content)} >>\nstream\n{content}\nendstream"
        kids.append(n)
        n += 2
    objs[2] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>"
    out, offs = "%PDF-1.4\n", {}
    for i in sorted(objs):
        offs[i] = len(out)
        out += f"{i} 0 obj\n{objs[i]}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {max(objs) + 1}\n0000000000 65535 f \n" + "".join(f"{offs[i]:010d} 00000 n \n" for i in sorted(objs))
    out += f"trailer\n<< /Size {max(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    data = out.encode("latin1")
    if encrypt:
        from pypdf import PdfReader, PdfWriter

        w = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
        w.encrypt(encrypt)
        buf = io.BytesIO()
        w.write(buf)
        return buf.getvalue()
    return data


def make_zip(entries: dict[str, bytes], deflate: bool = True) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED if deflate else zipfile.ZIP_STORED) as z:
        for name, data in entries.items():
            z.writestr(name, data)
    return out.getvalue()


def make_bomb_docx(mb: int = 5) -> bytes:
    return make_zip({"[Content_Types].xml": b"<Types/>", "word/document.xml": b"\0" * (mb * 1024 * 1024)})


def jpeg_with_exif(size=(64, 48), orientation: int = 6) -> bytes:
    from PIL import Image

    im = Image.new("RGB", size, (10, 120, 200))
    exif = Image.Exif()
    exif[0x010E] = "GPS 12.9716N 77.5946E Owner Asha"
    exif[0x0112] = orientation
    out = io.BytesIO()
    im.save(out, "JPEG", exif=exif)
    return out.getvalue()


def png_bytes(size=(40, 30), mode="RGB", colour=(200, 30, 30), noise: bool = False) -> bytes:
    from PIL import Image

    im = Image.new(mode, size, colour)
    if noise:
        im = Image.frombytes("RGB", size, os.urandom(size[0] * size[1] * 3))
    out = io.BytesIO()
    im.save(out, "PNG")
    return out.getvalue()


HAS_FFMPEG = shutil.which("ffmpeg") is not None and shutil.which("ffprobe") is not None


def make_video(tmp_path: Path, seconds: int = 6, audio: bool = True) -> bytes:
    """Three colours, two seconds each (two scene changes), with a sine tone: made by ffmpeg itself."""
    out = tmp_path / "v.mp4"
    third = seconds // 3
    graph = f"color=c=red:s=320x240:d={third}:r=10[a];color=c=green:s=320x240:d={third}:r=10[b];color=c=blue:s=320x240:d={third}:r=10[c];[a][b][c]concat=n=3:v=1:a=0,format=yuv420p[v]"
    cmd = ["ffmpeg", "-nostdin", "-v", "error", "-filter_complex", graph]
    if audio:
        cmd += ["-f", "lavfi", "-i", f"sine=frequency=440:duration={seconds}"]
    cmd += ["-map", "[v]"] + (["-map", "0:a"] if audio else []) + ["-c:v", "mpeg4", "-q:v", "5"] + (["-c:a", "aac", "-b:a", "32k"] if audio else []) + ["-y", str(out)]
    subprocess.run(cmd, check=True)
    return out.read_bytes()
